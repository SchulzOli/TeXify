"""
TeXify core — shared document/LaTeX logic.

Used by both the MCP server (server.py) and the web UI (web.py).
All files live under WORKSPACE_DIR (default /data), bind-mounted to the host.
"""

from __future__ import annotations

import os
import re
import subprocess
import unicodedata
from pathlib import Path

# --------------------------------------------------------------------------- #
# Config
# --------------------------------------------------------------------------- #

WORKSPACE = Path(os.environ.get("WORKSPACE_DIR", "/data")).resolve()
WORKSPACE.mkdir(parents=True, exist_ok=True)

PREVIEWABLE_SOURCES = {".pdf", ".docx", ".txt", ".md"}
TEXT_LIKE = {".tex", ".txt", ".md", ".log"}


# --------------------------------------------------------------------------- #
# Path / text helpers
# --------------------------------------------------------------------------- #


def safe_name(name: str) -> Path:
    """Resolve a user-supplied (relative) file name inside WORKSPACE.

    Prevents path traversal outside the workspace.
    """
    p = (WORKSPACE / name).resolve()
    if not str(p).startswith(str(WORKSPACE)):
        raise ValueError(f"path escapes workspace: {name!r}")
    return p


def latex_escape(text: str) -> str:
    """Escape LaTeX special characters in plain text."""
    repl = {
        "\\": r"\textbackslash{}",
        "&": r"\&",
        "%": r"\%",
        "$": r"\$",
        "#": r"\#",
        "_": r"\_",
        "{": r"\{",
        "}": r"\}",
        "~": r"\textasciitilde{}",
        "^": r"\textasciicircum{}",
    }
    return "".join(repl.get(ch, ch) for ch in text)


def slugify(s: str, default: str = "document") -> str:
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    s = re.sub(r"[^A-Za-z0-9\-_]+", "_", s).strip("_").lower()
    return s or default


def run(cmd: list[str], cwd: Path, timeout: int = 180) -> subprocess.CompletedProcess:
    return subprocess.run(
        cmd, cwd=str(cwd), capture_output=True, text=True, timeout=timeout
    )


# --------------------------------------------------------------------------- #
# Markdown-ish -> LaTeX body conversion
# --------------------------------------------------------------------------- #

_MD_BULLET = re.compile(r"^\s*[-*\u2022]\s+(.*)$")
_MD_H1 = re.compile(r"^#\s+(.*)$")
_MD_H2 = re.compile(r"^##\s+(.*)$")
_MD_H3 = re.compile(r"^###\s+(.*)$")


def _inline_md(text: str) -> str:
    esc = latex_escape(text)
    esc = re.sub(r"\*\*(.+?)\*\*", r"\\textbf{\1}", esc)
    esc = re.sub(r"(?<!\w)\*(.+?)\*(?!\w)", r"\\textit{\1}", esc)
    esc = re.sub(r"`(.+?)`", r"\\texttt{\1}", esc)
    return esc


def text_to_latex_body(raw: str) -> str:
    """Heuristically convert plain/markdown-ish text to a LaTeX body."""
    lines = raw.splitlines()
    out: list[str] = []
    itemize_open = False

    def close_itemize():
        nonlocal itemize_open
        if itemize_open:
            out.append("\\end{itemize}")
            itemize_open = False

    for line in lines:
        stripped = line.strip()
        if not stripped:
            close_itemize()
            out.append("")
            continue

        m = _MD_H1.match(stripped)
        if m:
            close_itemize()
            out.append(f"\\section*{{{_inline_md(m.group(1))}}}")
            continue
        m = _MD_H2.match(stripped)
        if m:
            close_itemize()
            out.append(f"\\subsection*{{{_inline_md(m.group(1))}}}")
            continue
        m = _MD_H3.match(stripped)
        if m:
            close_itemize()
            out.append(f"\\subsubsection*{{{_inline_md(m.group(1))}}}")
            continue

        m = _MD_BULLET.match(stripped)
        if m:
            if not itemize_open:
                out.append("\\begin{itemize}")
                itemize_open = True
            out.append(f"  \\item {_inline_md(m.group(1))}")
            continue

        close_itemize()
        out.append(_inline_md(stripped) + "\n")

    close_itemize()
    return "\n".join(out)


ARTICLE_PREAMBLE = r"""\documentclass[11pt,a4paper]{article}
\usepackage[utf8]{inputenc}
\usepackage[T1]{fontenc}
\usepackage[margin=2.4cm]{geometry}
\usepackage{enumitem}
\usepackage{parskip}
\usepackage[hidelinks]{hyperref}
\setlist{itemsep=2pt,topsep=3pt,parsep=0pt}
"""


def wrap_article(title: str, body: str) -> str:
    return (
        ARTICLE_PREAMBLE
        + f"\n\\title{{{latex_escape(title)}}}\n\\date{{}}\n\n"
        + "\\begin{document}\n\\maketitle\n\n"
        + body
        + "\n\n\\end{document}\n"
    )


# --------------------------------------------------------------------------- #
# Source extraction
# --------------------------------------------------------------------------- #


def extract_pdf_text(path: Path) -> str:
    r = run(["pdftotext", "-layout", str(path), "-"], cwd=WORKSPACE, timeout=120)
    if r.returncode == 0 and r.stdout.strip():
        return r.stdout
    try:
        import pymupdf  # type: ignore

        doc = pymupdf.open(str(path))
        return "\n".join(p.get_text() for p in doc)
    except Exception:
        pass
    raise RuntimeError(
        "could not extract text from PDF (no text layer?). "
        "Scanned PDFs need OCR first."
    )


def extract_docx_text(path: Path) -> str:
    try:
        import docx  # python-docx
    except Exception as exc:
        raise RuntimeError(f"python-docx unavailable: {exc}")
    d = docx.Document(str(path))
    chunks: list[str] = []
    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            chunks.append("")
            continue
        style = (p.style.name or "").lower()
        if style.startswith("heading 1"):
            chunks.append("# " + t)
        elif style.startswith("heading 2"):
            chunks.append("## " + t)
        elif style.startswith("heading"):
            chunks.append("### " + t)
        else:
            chunks.append(t)
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            chunks.append(" - " + " | ".join(cells))
    return "\n".join(chunks)


# --------------------------------------------------------------------------- #
# High-level operations
# --------------------------------------------------------------------------- #


def list_documents() -> dict:
    files = sorted(WORKSPACE.iterdir())
    groups: dict[str, list[str]] = {}
    for f in files:
        if f.is_file():
            groups.setdefault(f.suffix.lower() or "(none)", []).append(f.name)
    return {"workspace": str(WORKSPACE), "files": groups}


def convert_to_latex(source: str, output: str | None = None) -> dict:
    """Convert a PDF/DOCX/TXT/MD in the workspace to a .tex document."""
    src = safe_name(source)
    if not src.exists():
        return {"error": f"source not found: {source}"}

    ext = src.suffix.lower()
    try:
        if ext == ".pdf":
            raw = extract_pdf_text(src)
        elif ext == ".docx":
            raw = extract_docx_text(src)
        elif ext in {".txt", ".md"}:
            raw = src.read_text(encoding="utf-8", errors="replace")
        else:
            return {"error": f"unsupported source type: {ext}"}
    except RuntimeError as exc:
        return {"error": str(exc)}

    body = text_to_latex_body(raw)
    title = src.stem.replace("_", " ").replace("-", " ").title()
    tex = wrap_article(title, body)

    out_name = output or (src.stem + ".tex")
    if not out_name.endswith(".tex"):
        out_name += ".tex"
    dst = safe_name(out_name)
    dst.write_text(tex, encoding="utf-8")

    return {
        "created": dst.name,
        "source": src.name,
        "chars_extracted": len(raw),
        "preview": tex[:3000],
        "note": "Review and refine with edit_latex, then compile_pdf.",
    }


def create_latex(name: str, content: str, overwrite: bool = False) -> dict:
    if not name.endswith(".tex"):
        name += ".tex"
    p = safe_name(name)
    if p.exists() and not overwrite:
        return {"error": f"{p.name} exists; set overwrite=true to replace."}
    if "\\documentclass" not in content:
        return {
            "error": "content does not look like a full LaTeX document (no \\documentclass)."
        }
    p.write_text(content, encoding="utf-8")
    return {"created": p.name, "bytes": p.stat().st_size}


def edit_latex(name: str, old_string: str, new_string: str, replace_all: bool = False) -> dict:
    p = safe_name(name)
    if not p.exists():
        return {"error": f"not found: {name}"}
    text = p.read_text(encoding="utf-8")

    count = text.count(old_string)
    if count == 0:
        return {"error": "old_string not found; no edit applied."}
    if count > 1 and not replace_all:
        return {
            "error": f"old_string occurs {count} times; pass replace_all=true or add context."
        }

    new_text = (
        text.replace(old_string, new_string)
        if replace_all
        else text.replace(old_string, new_string, 1)
    )
    p.write_text(new_text, encoding="utf-8")

    idx = new_text.find(new_string)
    start = max(0, idx - 300)
    end = min(len(new_text), idx + len(new_string) + 300)
    return {
        "edited": p.name,
        "replacements": count if replace_all else 1,
        "context": new_text[start:end],
    }


def compile_pdf(name: str, passes: int = 2) -> dict:
    if not name.endswith(".tex"):
        name += ".tex"
    src = safe_name(name)
    if not src.exists():
        return {"error": f"not found: {name}"}

    pdf = src.with_suffix(".pdf")
    log_tail = ""
    for i in range(max(1, passes)):
        r = run(
            ["pdflatex", "-interaction=nonstopmode", "-halt-on-error", src.name],
            cwd=WORKSPACE,
            timeout=180,
        )
        if r.returncode != 0:
            log_file = src.with_suffix(".log")
            if log_file.exists():
                lines = log_file.read_text(encoding="utf-8", errors="replace").splitlines()
                err_idx = [n for n, ln in enumerate(lines) if ln.startswith("!")]
                if err_idx:
                    n = err_idx[-1]
                    log_tail = "\n".join(lines[n : n + 20])
                else:
                    log_tail = "\n".join(lines[-25:])
            else:
                log_tail = (r.stdout or "")[-2000:] + (r.stderr or "")[-500:]
            return {
                "success": False,
                "pass": i + 1,
                "error": "pdflatex failed",
                "log_tail": log_tail,
            }

    for aux in WORKSPACE.glob(src.stem + ".*"):
        if aux.suffix in {".aux", ".out", ".fdb_latexmk", ".fls", ".synctex.gz"}:
            aux.unlink(missing_ok=True)

    return {
        "success": True,
        "pdf": pdf.name,
        "tex": src.name,
        "pdf_bytes": pdf.stat().st_size if pdf.exists() else None,
    }


def delete_document(name: str) -> dict:
    p = safe_name(name)
    if not p.exists():
        return {"error": f"not found: {name}"}
    removed = [p.name]
    p.unlink()
    if p.suffix == ".tex":
        for aux in WORKSPACE.glob(p.stem + ".*"):
            if aux.suffix in {".aux", ".out", ".log", ".fdb_latexmk", ".fls", ".synctex.gz"}:
                removed.append(aux.name)
                aux.unlink(missing_ok=True)
    return {"removed": removed}


def save_upload(filename: str, data: bytes) -> dict:
    """Persist an uploaded file into the workspace (sanitised name)."""
    stem = slugify(Path(filename).stem)
    suffix = Path(filename).suffix.lower()
    if suffix not in PREVIEWABLE_SOURCES and suffix != ".tex":
        return {"error": f"unsupported upload type: {suffix or '(none)'}"}
    dst = safe_name(stem + suffix)
    dst.write_bytes(data)
    return {"saved": dst.name, "bytes": len(data)}
