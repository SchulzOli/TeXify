"""Import existing documents (PDF, DOCX, TXT) as Markdown.

Deterministic extraction only; restructuring the result into a specific
template (e.g. turning an old CV PDF into cv frontmatter) is an AI task:
``texify ai fill cv alter_lebenslauf.md``.
"""

from __future__ import annotations

import json
import re
from pathlib import Path

from . import TexifyError
from .pdftools import extract_text


def _docx_to_markdown(path: Path) -> str:
    try:
        import docx  # python-docx
    except ImportError:
        raise TexifyError("DOCX import needs python-docx:  pip install \"texify[docx]\"") from None
    d = docx.Document(str(path))
    out: list[str] = []
    for p in d.paragraphs:
        text = p.text.strip()
        style = (p.style.name or "").lower() if p.style is not None else ""
        if not text:
            out.append("")
        elif style.startswith("heading"):
            m = re.search(r"\d+", style)
            level = min(int(m.group()) if m else 2, 4)
            out += ["", "#" * level + " " + text, ""]
        elif "list" in style:
            out.append("- " + text)
        else:
            out += [text, ""]
    for table in d.tables:
        rows = [[c.text.strip().replace("\n", " ").replace("|", "/") for c in r.cells] for r in table.rows]
        if not rows:
            continue
        out += ["", "| " + " | ".join(rows[0]) + " |", "|" + "---|" * len(rows[0])]
        out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
        out.append("")
    return re.sub(r"\n{3,}", "\n\n", "\n".join(out)).strip() + "\n"


def _pdf_to_markdown(path: Path) -> str:
    text = extract_text(path)
    if not text.strip():
        raise TexifyError(f"{path.name}: no text layer found (scanned?). Run OCR first.")
    # A line continues the previous one only if that one did not end a sentence
    # and this one starts in lower case (hard-wrapped paragraph); otherwise it is
    # its own paragraph (headings, address lines, table rows ...).
    paragraphs: list[str] = []
    for line in (ln.strip() for ln in text.splitlines()):
        if not line:
            continue
        prev = paragraphs[-1] if paragraphs else ""
        if prev and not prev.endswith((".", ":", "!", "?")) and line[:1].islower():
            paragraphs[-1] = f"{prev} {line}"
        else:
            paragraphs.append(line)
    return "\n\n".join(paragraphs) + "\n"


def to_markdown(path: Path, template: str | None = None) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        body = _pdf_to_markdown(path)
    elif suffix == ".docx":
        body = _docx_to_markdown(path)
    elif suffix == ".md":
        return path.read_text(encoding="utf-8-sig")
    elif suffix == ".txt":
        body = path.read_text(encoding="utf-8-sig")
    else:
        raise TexifyError(f"Cannot import {suffix or 'files without extension'} (supported: pdf, docx, txt, md)")
    title = path.stem.replace("_", " ").replace("-", " ").strip()
    front = f"---\ntemplate: {template or 'article'}\ntitle: {json.dumps(title, ensure_ascii=False)}\n---\n\n"
    return front + body
